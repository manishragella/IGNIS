import os
import io
import json
import base64
import re
import requests
from typing import List, Optional
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from bs4 import BeautifulSoup
from dotenv import load_dotenv

# Import python-pptx and python-docx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Load environment variables from parent directory's .env.local
load_dotenv(dotenv_path="../.env.local")

# Set up Gemini
import google.generativeai as genai
api_key = os.environ.get("GEMINI_API_KEY")
if api_key:
    genai.configure(api_key=api_key)

app = FastAPI(title="Ignis Content & Export API")

# Configure CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/")
def read_root():
    return {
        "status": "Ignis AI Facilitator API is online 🔥",
        "version": "1.0.0",
        "author": "Ignis Team"
    }

# Parsers
def extract_text_from_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text = ""
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text += f"\n--- Page {i+1} ---\n{page_text}"
        return text
    except Exception as e:
        print(f"Error parsing PDF: {e}")
        return "[Error extracting text from PDF]"

def extract_text_from_docx(file_bytes: bytes) -> str:
    import docx2txt
    try:
        text = docx2txt.process(io.BytesIO(file_bytes))
        return text
    except Exception as e:
        print(f"Error parsing DOCX: {e}")
        return "[Error extracting text from DOCX]"

def fetch_and_extract_from_link(url: str) -> str:
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        res = requests.get(url, headers=headers, timeout=10)
        if res.ok:
            soup = BeautifulSoup(res.text, 'html.parser')
            # Remove scripts, styles
            for script in soup(["script", "style"]):
                script.decompose()
            # Extract heading and paragraph text
            elements = soup.find_all(['h1', 'h2', 'h3', 'p'])
            text_blocks = [el.get_text().strip() for el in elements if len(el.get_text().strip()) > 10]
            content = "\n".join(text_blocks[:20]) # Limit to first 20 blocks
            return f"\n--- URL: {url} ---\nTitle: {soup.title.string if soup.title else url}\n{content}"
        else:
            return f"\n--- URL: {url} ---\n[Failed to fetch content, HTTP Status {res.status_code}]"
    except Exception as e:
        return f"\n--- URL: {url} ---\n[Failed to fetch content due to error: {str(e)}]"

# Models
class ExportRequest(BaseModel):
    activity: dict
    includeIllustrations: bool = True

# Helper to generate illustrations
def generate_illustration(prompt: str) -> str:
    # Append high-fidelity Indian educational styling modifier
    style_modifier = ", vibrant Indian color palette, South Asian descent characters, local Indian context, flat vector educational illustration style"
    full_prompt = f"{prompt}{style_modifier}"
    print(f"Generating illustration with prompt: {full_prompt}")
    
    # Try keyless Pollinations AI API
    try:
        pollinations_url = f"https://image.pollinations.ai/p/{requests.utils.quote(full_prompt)}?width=1024&height=768&nologo=true"
        img_res = requests.get(pollinations_url, timeout=15)
        if img_res.ok:
            encoded_img = base64.b64encode(img_res.content).decode('utf-8')
            return f"data:image/jpeg;base64,{encoded_img}"
        else:
            print(f"Pollinations AI failed with status: {img_res.status_code}")
    except Exception as e:
        print(f"Error calling Pollinations AI: {e}")
    
    # Fallback placeholder if failed
    return ""

@app.post("/api/generate-activity")
async def generate_activity(
    grade: str = Form(...),
    topic: str = Form(...),
    languageFocus: str = Form(...),  # Expecting comma separated or JSON string
    lifeSkillFocus: str = Form(...),  # Expecting comma separated or JSON string
    includeIllustrations: str = Form("false"),
    links: str = Form("[]"),
    files: Optional[List[UploadFile]] = File(None)
):
    try:
        # Resolve array parameter lists
        try:
            lang_focus_list = json.loads(languageFocus)
            if not isinstance(lang_focus_list, list):
                lang_focus_list = [str(lang_focus_list)]
        except:
            lang_focus_list = [x.strip() for x in languageFocus.split(",") if x.strip()]

        try:
            life_skill_list = json.loads(lifeSkillFocus)
            if not isinstance(life_skill_list, list):
                life_skill_list = [str(life_skill_list)]
        except:
            life_skill_list = [x.strip() for x in lifeSkillFocus.split(",") if x.strip()]

        try:
            links_list = json.loads(links)
        except:
            links_list = [x.strip() for x in links.split(",") if x.strip()]

        include_illustrations = includeIllustrations.lower() == "true"

        # Document Parsing
        context_text = ""
        citations = []
        source_id = 1

        if files:
            for upload_file in files:
                file_bytes = await upload_file.read()
                file_text = ""
                if upload_file.filename.endswith('.pdf'):
                    file_text = extract_text_from_pdf(file_bytes)
                elif upload_file.filename.endswith(('.docx', '.doc')):
                    file_text = extract_text_from_docx(file_bytes)
                else:
                    try:
                        file_text = file_bytes.decode('utf-8')
                    except:
                        file_text = "[Unsupported file format]"
                
                if file_text and len(file_text) > 10:
                    context_text += f"\nSource [{source_id}] (File): {upload_file.filename}\n{file_text[:8000]}\n" # Cap at 8k chars per file
                    citations.append({
                        "id": source_id,
                        "source": upload_file.filename,
                        "type": "File",
                        "summary": f"Provided background ideas and guidelines about '{topic}'."
                    })
                    source_id += 1
        
        if links_list:
            for link in links_list:
                if link.startswith('http'):
                    link_text = fetch_and_extract_from_link(link)
                    if link_text:
                        context_text += f"\nSource [{source_id}] (Link): {link}\n{link_text}\n"
                        citations.append({
                            "id": source_id,
                            "source": link,
                            "type": "Link",
                            "summary": f"Retrieved web documentation reference."
                        })
                        source_id += 1

        # Hyper-Personalization details based on Grade Level
        # Adapt cognitive capability, sentence structure, and task complexity
        grade_str = grade.lower()
        if "class 3" in grade_str or "grade 3" in grade_str:
            personalization_guideline = """
            - Grade Adaptation: Grade 3 / Class 3 focus. Use simple vocabulary suitable for 8-9 year olds. Keep sentences short (under 10-12 words). Focus on concrete examples, visual descriptions, and interactive small-group games.
            - Life Skill Focus: Focus on basic cooperation, empathy, and active listening.
            - Presentation Task: Keep it very simple (e.g. show a drawing, share one sentence, roleplay a 30-second scene in pairs).
            """
        elif "class 5" in grade_str or "grade 5" in grade_str:
            personalization_guideline = """
            - Grade Adaptation: Grade 5 / Class 5 focus. Introduce basic abstract thinking. Suitable for 10-11 year olds. Use moderately challenging vocabulary. Sentences can be complex. Include cooperative group structures with roles, analytical questions, and structured problem-solving.
            - Life Skill Focus: Focus on structured teamwork, critical decision making, and reflective speaking.
            - Presentation Task: Group poster presentations, presenting structured findings, or pitching a collaborative solution.
            """
        else:
            personalization_guideline = f"""
            - Grade Adaptation: General level adaptation for {grade}. Vocabulary and activities should be age-appropriate.
            """

        # System and User Prompts
        system_prompt = f"""
        You are an expert educational facilitator for Ignis, an activity-based learning platform tailored specifically for the Indian educational context.
        
        Generate a comprehensive, hyper-personalized, step-by-step classroom activity worksheet and gamified self-learning widgets.
        
        CRITICAL - INLINE CITATIONS & REFERENCES:
        We have parsed reference source materials from the user. These are numbered as Source [1], Source [2], etc.
        - Whenever you use factual claims, concepts, or terms directly sourced from these reference materials, you MUST append a superscript-style citation, e.g., '[1]', '[2]' in your text content (for example: "Rainwater harvesting has been practiced in India for thousands of years [1]...").
        - Keep citations clean, numerical, and map them precisely. Only reference valid, existing source numbers. Do not invent source IDs.
        
        CRITICAL - INDIAN CONTEXT ENFORCEMENT:
        You MUST design this activity with rich Indian cultural references, local settings, and relatable naming conventions:
        - Settings: Set activities in diverse Indian environments (e.g., Pune, a village panchayat near Jaipur, a municipal school in Kochi).
        - Names: Use common Indian names for student roles (e.g. Aarav, Priya, Rohan, Sunita, Kabir, Meera, Arjun, Ananya).
        - Local Contexts: Reference relatable Indian items or issues (e.g. monsoons, regional crops, traditional crafts, local community helpers).
        
        Generate a valid JSON object matching the following structure EXACTLY:
        {{
          "title": "Engaging activity title reflecting the Indian theme",
          "objectives": ["Objective 1", "Objective 2", "Objective 3"],
          "estimatedTime": "e.g. 45 minutes",
          "materials": ["Material 1", "Material 2"],
          "teacherInstructions": "Detailed overview instructions for the teacher to facilitate, including any inline citations like [1] where appropriate.",
          "presentationTask": "Description of what students will present or create as final output",
          "slides": [
            {{
              "stepNumber": 1,
              "stepTitle": "Title of Step 1 (e.g., Warm Up / Introduction)",
              "stepContent": "Clear, engaging explanation of this step tailored for the classroom. Can include citations like [1].",
              "imagePrompt": "A highly detailed, child-friendly, colorful vector illustration prompt describing an educational scene relevant to this step. Must feature Indian students or teachers in school uniforms, South Asian descent, vibrant Indian colors, no text, clean educational vector style."
            }},
            {{
              "stepNumber": 2,
              "stepTitle": "Title of Step 2 (e.g., Group Activity / Main Task)",
              "stepContent": "Clear instructions for what the students will do next in their groups.",
              "imagePrompt": "Detailed vector art illustration prompt representing this step. Features Indian children collaborating, local context, clean white background, kid-friendly vector."
            }},
            {{
              "stepNumber": 3,
              "stepTitle": "Title of Step 3 (e.g., Student Presentation / Shared Reflection)",
              "stepContent": "How students will present their ideas or reflect on the life skill.",
              "imagePrompt": "Detailed vector art illustration prompt showing Indian classroom presentations or sharing ideas."
            }}
          ],
          "worksheet": [
            {{
              "heading": "Introduction & Context",
              "content": "Student-facing warm up written in accessible grade-level language. Append citations like [1] where appropriate.",
              "type": "instructions"
            }},
            {{
              "heading": "Our Main Action Task",
              "content": "Detailed step-by-step instructions written directly for the students. Include inline citations.",
              "type": "activity"
            }},
            {{
              "heading": "Reflective Discussion",
              "content": "Questions for the student groups to answer together",
              "type": "questions"
            }}
          ],
          "rubric": [
            {{
              "criterion": "Language Application (Focus: {', '.join(lang_focus_list)})",
              "excellent": "Excellent description",
              "good": "Good description",
              "developing": "Developing description",
              "beginning": "Beginning description"
            }},
            {{
              "criterion": "Life Skills (Focus: {', '.join(life_skill_list)})",
              "excellent": "Excellent description",
              "good": "Good description",
              "developing": "Developing description",
              "beginning": "Beginning description"
            }}
          ],
          "citations": [
            {{
              "id": 1,
              "source": "Name of Source 1",
              "summary": "1-2 sentence summary of what factual information was extracted and applied in this activity."
            }}
          ],
          "creativity_rationale": {{
            "factual_basis": "Explain in detail which factual parts of the activity were directly extracted from the source files or links (citing [1], [2], etc.). Keep it professional, monochromatic, and educational.",
            "creative_adaptations": "Explain the creative choices made (e.g. why we selected Pune as the location, why Rohan and Priya are water heroes, why this specific game/sorting metaphor was chosen for this class level) to increase engagement."
          }},
          "gamified_activities": {{
            "sorting": {{
              "title": "Sorting Challenge",
              "description": "Drag or click the scenarios to place them into the correct categories!",
              "categories": ["Category A (e.g. Do's)", "Category B (e.g. Don'ts)"],
              "scenarios": [
                {{
                  "id": "s1",
                  "text": "Scenario 1 (e.g. Turning off the tap)",
                  "correctCategory": "Category A"
                }},
                {{
                  "id": "s2",
                  "text": "Scenario 2 (e.g. Leaving water running)",
                  "correctCategory": "Category B"
                }}
              ]
            }},
            "blanks": {{
              "title": "Vocabulary Fill-in-the-Blanks",
              "description": "Fill in the missing words to complete the sentences correctly!",
              "sentences": [
                {{
                  "text_before": "Text before the blank (e.g. The Pune school has a water )",
                  "blank_key": "tank",
                  "text_after": " to save monsoons."
                }}
              ]
            }}
          }}
        }}
        """

        reference_context_str = ""
        if context_text:
            reference_context_str = f"Reference Context from uploaded sources:\n{context_text}"

        user_prompt = f"""
        Generate the activity with these parameters:
        - Grade Level: {grade}
        - Topic: {topic}
        - Target Language Focus: {', '.join(lang_focus_list)}
        - Target Life Skill Focus: {', '.join(life_skill_list)}
        
        Personalization Guidelines:
        {personalization_guideline}
        
        {reference_context_str}
        
        If source context is provided, synthesize it logically into the activity, include exact inline superscript citations (like [1], [2]) mapping back to the Source IDs, list citations in the citations array, write the AI insights, and construct the two customized playable games in 'gamified_activities'.
        Return ONLY the raw JSON block. No markdown markers (like ```json), no trailing text.
        """

        # Generate using Gemini
        if not api_key:
            raise HTTPException(status_code=500, detail="Gemini API Key is not configured. Please check your environmental variables.")

        model = genai.GenerativeModel('gemini-2.5-flash')
        combined_prompt = f"{system_prompt}\n\nUSER REQUEST PARAMETERS AND CONTEXT:\n{user_prompt}"
        response = model.generate_content(
            contents=combined_prompt,
            generation_config=genai.types.GenerationConfig(
                temperature=0.7
            )
        )

        res_text = response.text.strip()
        # Clean JSON wrappers if present
        res_text = re.sub(r"^```json\s*", "", res_text)
        res_text = re.sub(r"\s*```$", "", res_text)
        
        activity_data = json.loads(res_text)

        # Synthesize citations with AI summary
        generated_citations = activity_data.get("citations", [])
        if citations:
            citations_map = {c["id"]: c for c in citations}
            merged_citations = []
            for gen_cit in generated_citations:
                cit_id = gen_cit.get("id")
                if cit_id in citations_map:
                    merged_citations.append({
                        "id": cit_id,
                        "source": citations_map[cit_id]["source"],
                        "type": citations_map[cit_id]["type"],
                        "summary": gen_cit.get("summary", citations_map[cit_id]["summary"])
                    })
                else:
                    merged_citations.append(gen_cit)
            if not merged_citations:
                merged_citations = citations
            activity_data["citations"] = merged_citations
        else:
            activity_data["citations"] = generated_citations

        # Generate illustrations if enabled
        if include_illustrations and "slides" in activity_data:
            print("Illustration toggle is ON. Painting Indian-style visuals for each step...")
            for slide in activity_data["slides"]:
                img_prompt = slide.get("imagePrompt", f"Illustration for {topic}")
                base64_img = generate_illustration(img_prompt)
                if base64_img:
                    slide["imageUrl"] = base64_img
                else:
                    # Soft fallback if fails
                    slide["imageUrl"] = "/images/textbook-illustration-fallback.svg"
        else:
            # Set fallback/empty image urls if toggled off
            if "slides" in activity_data:
                for slide in activity_data["slides"]:
                    slide["imageUrl"] = ""

        return {"activity": activity_data}

    except Exception as e:
        print(f"Exception during generation: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to generate activity: {str(e)}")


@app.post("/api/export-ppt")
async def export_ppt(req: ExportRequest):
    try:
        activity = req.activity
        include_illustrations = req.includeIllustrations
        
        prs = Presentation()
        # Set standard 16:9 widescreen layout
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        # Color Palette - Elegant Dark Slate Blue & Orange/Marigold
        dark_bg = RGBColor(15, 23, 42)      # Slate 900
        gold_accent = RGBColor(245, 158, 11) # Amber 500 (Marigold)
        white_text = RGBColor(255, 255, 255)
        slate_text = RGBColor(51, 65, 85)    # Slate 700
        gray_bg = RGBColor(248, 250, 252)    # Slate 50
        
        # ----------------------------------------------------
        # Slide 1: TITLE SLIDE (Ignis Theme Dark)
        # ----------------------------------------------------
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        
        # Background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = dark_bg
        
        # Title text box
        tx_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.0))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = "🔥 IGNIS ACTIVITY DECK"
        p.alignment = PP_ALIGN.LEFT
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = gold_accent
        
        p2 = tf.add_paragraph()
        p2.text = activity.get("title", "Classroom Activity").upper()
        p2.alignment = PP_ALIGN.LEFT
        p2.font.size = Pt(44)
        p2.font.bold = True
        p2.font.color.rgb = white_text
        p2.space_before = Pt(12)
        
        p3 = tf.add_paragraph()
        p3.text = f"Time: {activity.get('estimatedTime', '45 mins')}   |   Citations Appended"
        p3.alignment = PP_ALIGN.LEFT
        p3.font.size = Pt(16)
        p3.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
        p3.space_before = Pt(20)
        
        # ----------------------------------------------------
        # Slide 2: OBJECTIVES & MATERIALS (Slate Light)
        # ----------------------------------------------------
        slide2 = prs.slides.add_slide(blank_layout)
        bg2 = slide2.background
        fill2 = bg2.fill
        fill2.solid()
        fill2.fore_color.rgb = gray_bg
        
        # Left Panel: Learning Objectives
        left_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(5.5), Inches(5.8))
        tf_left = left_box.text_frame
        tf_left.word_wrap = True
        
        p_obj_title = tf_left.paragraphs[0]
        p_obj_title.text = "🎯 Learning Objectives"
        p_obj_title.font.size = Pt(24)
        p_obj_title.font.bold = True
        p_obj_title.font.color.rgb = dark_bg
        p_obj_title.space_after = Pt(14)
        
        for obj in activity.get("objectives", ["Understand key concepts"]):
            p_item = tf_left.add_paragraph()
            p_item.text = f"• {obj}"
            p_item.font.size = Pt(16)
            p_item.font.color.rgb = slate_text
            p_item.space_before = Pt(8)
            
        # Right Panel: Materials & General Guide
        right_box = slide2.shapes.add_textbox(Inches(7.0), Inches(0.8), Inches(5.5), Inches(5.8))
        tf_right = right_box.text_frame
        tf_right.word_wrap = True
        
        p_mat_title = tf_right.paragraphs[0]
        p_mat_title.text = "📦 Materials Needed"
        p_mat_title.font.size = Pt(24)
        p_mat_title.font.bold = True
        p_mat_title.font.color.rgb = dark_bg
        p_mat_title.space_after = Pt(14)
        
        for mat in activity.get("materials", ["Notebooks", "Pens"]):
            p_item = tf_right.add_paragraph()
            p_item.text = f"• {mat}"
            p_item.font.size = Pt(16)
            p_item.font.color.rgb = slate_text
            p_item.space_before = Pt(8)
            
        # ----------------------------------------------------
        # Slides 3+: STEP-BY-STEP ACTIVITY STEPS (1-to-1 Split)
        # ----------------------------------------------------
        for step in activity.get("slides", []):
            slide_step = prs.slides.add_slide(blank_layout)
            
            # Step background (light)
            bg_s = slide_step.background
            fill_s = bg_s.fill
            fill_s.solid()
            fill_s.fore_color.rgb = white_text
            
            step_num = step.get("stepNumber", 1)
            step_title = re.sub(r'\[\d+\]', '', step.get("stepTitle", f"Step {step_num}")).strip()
            step_content = re.sub(r'\[\d+\]', '', step.get("stepContent", "")).strip()
            img_url = step.get("imageUrl", "")
            
            # Check if we should insert the illustration
            has_image = include_illustrations and img_url and img_url.startswith("data:image")
            
            # Slide Title Header
            header_box = slide_step.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.8))
            tf_hdr = header_box.text_frame
            tf_hdr.word_wrap = True
            p_hdr = tf_hdr.paragraphs[0]
            p_hdr.text = f"Step {step_num}: {step_title}".upper()
            p_hdr.font.size = Pt(20)
            p_hdr.font.bold = True
            p_hdr.font.color.rgb = gold_accent
            
            if has_image:
                # SPLIT SCREEN DESIGN: Left 50% Text, Right 50% Illustration
                # Left Text Box (50% Width)
                text_box = slide_step.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.5), Inches(5.2))
                tf_text = text_box.text_frame
                tf_text.word_wrap = True
                
                # We can chunk/paragraph the step content to make it clean
                p_body = tf_text.paragraphs[0]
                p_body.text = step_content
                p_body.font.size = Pt(18)
                p_body.font.color.rgb = slate_text
                p_body.space_before = Pt(6)
                p_body.line_spacing = 1.2
                
                # Right Image Box (50% Width)
                try:
                    # Decode base64 image data
                    img_data = base64.b64decode(img_url.split(",")[1])
                    img_io = io.BytesIO(img_data)
                    # Insert in 50% right area: Left=7.0", Top=1.3", Width=5.5", Height=4.8"
                    slide_step.shapes.add_picture(img_io, Inches(7.0), Inches(1.3), Inches(5.5), Inches(4.8))
                except Exception as ex:
                    print(f"Error inserting slide image: {ex}")
                    # Render a placeholder rectangle or icon if failed
                    placeholder_box = slide_step.shapes.add_textbox(Inches(7.0), Inches(1.3), Inches(5.5), Inches(4.8))
                    tf_p = placeholder_box.text_frame
                    p_p = tf_p.paragraphs[0]
                    p_p.text = "[Illustration Fallback: " + step.get("imagePrompt", "")[:50] + "...]"
                    p_p.font.size = Pt(14)
                    p_p.font.color.rgb = RGBColor(150, 150, 150)
            else:
                # FULL SCREEN TEXT LAYOUT
                text_box = slide_step.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.733), Inches(5.2))
                tf_text = text_box.text_frame
                tf_text.word_wrap = True
                
                p_body = tf_text.paragraphs[0]
                p_body.text = step_content
                p_body.font.size = Pt(20)
                p_body.font.color.rgb = slate_text
                p_body.space_before = Pt(6)
                p_body.line_spacing = 1.3

        # ----------------------------------------------------
        # Slide Last: CITATIONS & SOURCES (Dark Slate)
        # ----------------------------------------------------
        slide_last = prs.slides.add_slide(blank_layout)
        bg_l = slide_last.background
        fill_l = bg_l.fill
        fill_l.solid()
        fill_l.fore_color.rgb = dark_bg
        
        tx_box_last = slide_last.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.8))
        tf_last = tx_box_last.text_frame
        tf_last.word_wrap = True
        
        p_last_hdr = tf_last.paragraphs[0]
        p_last_hdr.text = "📚 Citations & Source Materials"
        p_last_hdr.font.size = Pt(26)
        p_last_hdr.font.bold = True
        p_last_hdr.font.color.rgb = gold_accent
        p_last_hdr.space_after = Pt(20)
        
        cits = activity.get("citations", [])
        if cits:
            for c in cits:
                p_c = tf_last.add_paragraph()
                p_c.text = f"📄 {c.get('source', 'Reference')}"
                p_c.font.size = Pt(18)
                p_c.font.bold = True
                p_c.font.color.rgb = white_text
                p_c.space_before = Pt(12)
                
                p_c_desc = tf_last.add_paragraph()
                p_c_desc.text = c.get('summary', 'Retrieved educational data source.')
                p_c_desc.font.size = Pt(14)
                p_c_desc.font.color.rgb = RGBColor(203, 213, 225) # Slate 300
                p_c_desc.space_before = Pt(2)
        else:
            p_empty = tf_last.add_paragraph()
            p_empty.text = "All content was synthesized by the Ignis AI Engine based on standard Class-level curricula guidelines."
            p_empty.font.size = Pt(16)
            p_empty.font.color.rgb = RGBColor(203, 213, 225)
            p_empty.space_before = Pt(20)

        # Output to Stream
        ppt_stream = io.BytesIO()
        prs.save(ppt_stream)
        ppt_stream.seek(0)
        
        return StreamingResponse(
            ppt_stream,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f"attachment; filename={activity.get('title', 'activity').replace(' ', '_')}_deck.pptx"}
        )

    except Exception as e:
        print(f"PPT Export Exception: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to generate PPT presentation: {str(e)}")


@app.post("/api/export-docx")
async def export_docx(req: ExportRequest):
    try:
        activity = req.activity
        include_illustrations = req.includeIllustrations
        
        doc = Document()
        
        # Color definitions for Word styling
        dark_blue = DocRGBColor(30, 41, 59)     # Slate 800
        gold_accent = DocRGBColor(217, 119, 6)   # Amber 600
        slate_gray = DocRGBColor(100, 116, 139)  # Slate 500
        
        # Title of Document
        title = doc.add_heading(level=0)
        run_title = title.add_run(activity.get("title", "Ignis Learning Activity Worksheet").upper())
        run_title.font.name = "Arial"
        run_title.font.size = DocPt(24)
        run_title.font.bold = True
        run_title.font.color.rgb = dark_blue
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Meta info
        meta_p = doc.add_paragraph()
        run_time = meta_p.add_run(f"Estimated Time: {activity.get('estimatedTime', '45 minutes')}  |  Ignis AI Facilitator")
        run_time.font.name = "Arial"
        run_time.font.size = DocPt(10)
        run_time.font.italic = True
        run_time.font.color.rgb = slate_gray
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph().paragraph_format.space_after = DocPt(10)
        
        # Section 1: Objectives
        h1 = doc.add_heading(level=1)
        r1 = h1.add_run("1. LEARNING OBJECTIVES")
        r1.font.bold = True
        r1.font.size = DocPt(16)
        r1.font.color.rgb = gold_accent
        
        for obj in activity.get("objectives", []):
            doc.add_paragraph(obj, style='List Bullet')
            
        # Section 2: Materials Required
        h2 = doc.add_heading(level=1)
        r2 = h2.add_run("2. MATERIALS REQUIRED")
        r2.font.bold = True
        r2.font.size = DocPt(16)
        r2.font.color.rgb = gold_accent
        
        for mat in activity.get("materials", []):
            doc.add_paragraph(mat, style='List Bullet')
            
        # Section 3: Teacher's Implementation Guide
        h3 = doc.add_heading(level=1)
        r3 = h3.add_run("3. TEACHER'S IMPLEMENTATION GUIDE")
        r3.font.bold = True
        r3.font.size = DocPt(16)
        r3.font.color.rgb = gold_accent
        
        guide_p = doc.add_paragraph()
        r_guide = guide_p.add_run(activity.get("teacherInstructions", "No specific teacher instructions provided."))
        r_guide.font.size = DocPt(11)
        r_guide.font.name = "Georgia"
        
        # Section 4: Step-by-Step Activities (Slide mapping)
        h4 = doc.add_heading(level=1)
        r4 = h4.add_run("4. SEQUENTIAL LESSON PLAN STEPS")
        r4.font.bold = True
        r4.font.size = DocPt(16)
        r4.font.color.rgb = gold_accent
        
        for step in activity.get("slides", []):
            step_num = step.get("stepNumber", 1)
            step_title = step.get("stepTitle", f"Step {step_num}")
            step_content = step.get("stepContent", "")
            img_url = step.get("imageUrl", "")
            
            # Step Heading
            step_h = doc.add_heading(level=2)
            r_step_h = step_h.add_run(f"Step {step_num}: {step_title}")
            r_step_h.font.bold = True
            r_step_h.font.size = DocPt(13)
            r_step_h.font.color.rgb = dark_blue
            
            # Step Content
            step_p = doc.add_paragraph()
            r_step_c = step_p.add_run(step_content)
            r_step_c.font.size = DocPt(11)
            r_step_c.font.name = "Arial"
            
            # Step Image (if present and toggled)
            if include_illustrations and img_url and img_url.startswith("data:image"):
                try:
                    img_data = base64.b64decode(img_url.split(",")[1])
                    img_io = io.BytesIO(img_data)
                    # Embed inline image centered, width 4.5 inches
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_run = img_p.add_run()
                    img_run.add_picture(img_io, width=DocInches(4.5))
                    
                    # Caption
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_cap = cap_p.add_run(f"Figure {step_num}: Illustration for {step_title}")
                    r_cap.font.size = DocPt(9)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = slate_gray
                except Exception as ex:
                    print(f"Error embedding word doc image: {ex}")
        
        # Section 5: Student Worksheet Sections
        h5 = doc.add_heading(level=1)
        r5 = h5.add_run("5. CLASSROOM STUDENT WORKSHEET")
        r5.font.bold = True
        r5.font.size = DocPt(16)
        r5.font.color.rgb = gold_accent
        
        for w_sec in activity.get("worksheet", []):
            sec_h = doc.add_heading(level=2)
            r_sec_h = sec_h.add_run(w_sec.get("heading", "Worksheet Part"))
            r_sec_h.font.bold = True
            r_sec_h.font.size = DocPt(12)
            
            sec_p = doc.add_paragraph()
            r_sec_c = sec_p.add_run(w_sec.get("content", ""))
            r_sec_c.font.size = DocPt(11)
            r_sec_c.font.name = "Arial"
            
        # Section 6: Final Presentation Task
        if activity.get("presentationTask"):
            h6 = doc.add_heading(level=1)
            r6 = h6.add_run("6. PRESENTATION & SUBMISSION TASK")
            r6.font.bold = True
            r6.font.size = DocPt(16)
            r6.font.color.rgb = gold_accent
            
            pres_p = doc.add_paragraph()
            r_pres = pres_p.add_run(activity.get("presentationTask"))
            r_pres.font.size = DocPt(11)
            r_pres.font.name = "Arial"
            
        # Section 7: Assessment Rubric Table
        if activity.get("rubric"):
            h7 = doc.add_heading(level=1)
            r7 = h7.add_run("7. ASSESSMENT RUBRIC")
            r7.font.bold = True
            r7.font.size = DocPt(16)
            r7.font.color.rgb = gold_accent
            
            # Add table
            table = doc.add_table(rows=1, cols=5)
            table.style = 'Light Shading Accent 1'
            hdr_cells = table.rows[0].cells
            hdr_titles = ["Criterion", "Excellent (4)", "Good (3)", "Developing (2)", "Beginning (1)"]
            for idx, title_text in enumerate(hdr_titles):
                hdr_cells[idx].text = title_text
                
            for rub in activity.get("rubric", []):
                row_cells = table.add_row().cells
                row_cells[0].text = rub.get("criterion", "Criterion")
                row_cells[1].text = rub.get("excellent", "")
                row_cells[2].text = rub.get("good", "")
                row_cells[3].text = rub.get("developing", "")
                row_cells[4].text = rub.get("beginning", "")
        
        # Section 8: Citations
        cits = activity.get("citations", [])
        if cits:
            doc.add_page_break()
            h8 = doc.add_heading(level=1)
            r8 = h8.add_run("8. CITATIONS & DATA REFERENCES")
            r8.font.bold = True
            r8.font.size = DocPt(16)
            r8.font.color.rgb = gold_accent
            
            for c in cits:
                cite_p = doc.add_paragraph()
                r_cite_t = cite_p.add_run(f"• {c.get('source', 'Source Reference')} ({c.get('type', 'Web')})\n")
                r_cite_t.font.bold = True
                
                r_cite_s = cite_p.add_run(c.get('summary', 'Retrieved reference details.'))
                r_cite_s.font.size = DocPt(10)
                r_cite_s.font.italic = True
                r_cite_s.font.color.rgb = slate_gray

        # Section 9: Creativity Rationale (AI Insights)
        rationale = activity.get("creativity_rationale")
        if rationale:
            doc.add_page_break()
            h9 = doc.add_heading(level=1)
            r9 = h9.add_run("9. AI INSIGHTS & PEDAGOGICAL RATIONALE")
            r9.font.bold = True
            r9.font.size = DocPt(16)
            r9.font.color.rgb = gold_accent
            
            fact_h = doc.add_heading(level=2)
            r_fact_h = fact_h.add_run("Factual Basis & Contextual Tracing")
            r_fact_h.font.bold = True
            r_fact_h.font.size = DocPt(12)
            
            fact_p = doc.add_paragraph()
            r_fact_c = fact_p.add_run(rationale.get("factual_basis", ""))
            r_fact_c.font.size = DocPt(10.5)
            r_fact_c.font.name = "Arial"
            
            adapt_h = doc.add_heading(level=2)
            r_adapt_h = adapt_h.add_run("Creative Adaptations & Pedagogical Choices")
            r_adapt_h.font.bold = True
            r_adapt_h.font.size = DocPt(12)
            
            adapt_p = doc.add_paragraph()
            r_adapt_c = adapt_p.add_run(rationale.get("creative_adaptations", ""))
            r_adapt_c.font.size = DocPt(10.5)
            r_adapt_c.font.name = "Arial"
        docx_stream = io.BytesIO()
        doc.save(docx_stream)
        docx_stream.seek(0)
        
        return StreamingResponse(
            docx_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={activity.get('title', 'activity').replace(' ', '_')}_worksheet.docx"}
        )

    except Exception as e:
        print(f"DOCX Export Exception: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to generate Word document: {str(e)}")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
